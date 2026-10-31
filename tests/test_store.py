from collections import Counter

from docx import Document

import ingest
import settings
import store
from ingest import ingest_dir, split_chunks


def _fresh(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "CHROMA", tmp_path / "chroma")
    store.reset()


def test_upsert_skips_same_checksum(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    chunks = split_chunks("Цех имеет номер 4. Мастер цеха Иванов.", "charter.txt")
    store.upsert_chunks(chunks)
    store.upsert_chunks(chunks)
    assert store.collection().count() == 1


def test_reingest_applies_overlap_change(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    monkeypatch.setattr(store, "_embed", lambda texts: [[0.1, 0.2, 0.3, 0.4] for _ in texts])
    monkeypatch.setattr(ingest, "count_tokens", lambda text: max(1, len((text or "").split())))
    ingest._word_tokens.cache_clear()
    text = "\n".join(
        f"строка {number} " + " ".join(f"слово{number}_{index}" for index in range(8))
        for number in range(40)
    )
    state = {"tokens": 250, "overlap": 40, "intro": 120}
    monkeypatch.setattr(ingest, "runtime", lambda: state)
    first = split_chunks(text, "doc.txt")
    store.upsert_chunks(first)
    state["overlap"] = 0
    second = split_chunks(text, "doc.txt")
    assert len(first) == len(second) > 1
    assert [chunk["text"] for chunk in first] != [chunk["text"] for chunk in second]
    store.upsert_chunks(second)
    _, texts, metas = store._rows(store.collection())
    stored = [body for _, body in sorted(zip(metas, texts), key=lambda item: item[0]["chunk_index"])]
    assert stored == [chunk["text"] for chunk in second]
    ingest._word_tokens.cache_clear()


def test_upsert_replaces_changed_document(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    store.upsert_chunks(split_chunks("Насос синий.", "pump.md"))
    store.upsert_chunks(split_chunks("Насос зелёный.", "pump.md"))
    assert store._rows(store.collection())[1] == ["Насос зелёный."]


def test_chunks_for_document_orders_by_index(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    monkeypatch.setattr(store, "_embed", lambda texts: [[0.1, 0.2, 0.3, 0.4] for _ in texts])
    store.upsert_chunks([
        {"text": "второй", "document_name": "note.txt", "chunk_index": 1, "checksum": "a", "section": "B"},
        {"text": "первый", "document_name": "note.txt", "chunk_index": 0, "checksum": "a", "section": "A"},
        {"text": "чужой", "document_name": "other.txt", "chunk_index": 0, "checksum": "b", "section": ""},
    ])
    rows = store.chunks_for_document("note.txt")
    assert [row["text"] for row in rows] == ["первый", "второй"]
    assert [row["section"] for row in rows] == ["A", "B"]
    assert [row["chunk_index"] for row in rows] == [0, 1]


def test_empty_search_returns_empty(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "CHROMA", tmp_path / "empty")
    store.reset()
    assert store.search_chunks("номер цеха") == []


def test_search_finds_document(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    store.upsert_chunks(split_chunks("Насос имеет обозначение Н-12. Рабочее давление равно 6 бар.", "pump.md"))
    found = store.search_chunks("давление насоса")
    assert found[0]["document_name"] == "pump.md"
    assert "6 бар" in found[0]["text"]


def test_search_keeps_to_document_named_in_question(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    store.upsert_chunks(split_chunks("Система АЛЬФА передаёт данные по протоколу HTTP.", "a.txt"))
    store.upsert_chunks(split_chunks("Система БЕТА передаёт данные по протоколу FTP.", "b.txt"))
    found = store.search_chunks("По какому протоколу передаёт данные БЕТА?")
    assert {chunk["document_name"] for chunk in found} == {"b.txt"}


def test_search_keeps_to_numbered_filename_in_question(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    same = "В разделе Назначение программы должны быть указаны сведения о назначении программы."
    store.upsert_chunks(split_chunks(same, "STD_1.2-99.txt"))
    store.upsert_chunks(split_chunks(same, "STD_1.3-99.txt"))
    found = store.search_chunks("Что указывают в разделе Назначение программы по STD 1.2-99?")
    assert {chunk["document_name"] for chunk in found} == {"STD_1.2-99.txt"}


def test_search_finds_exact_code(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    lines = [f"Поток X_{number}.1 передаёт заявки номер {number}." for number in range(1, 9)]
    lines.append("Поток X_9.2 передаёт годовые отчёты.")
    store.upsert_chunks(split_chunks("\n".join(lines * 20), "flows.txt"))
    found = store.search_chunks("Что передаёт поток X_9.2?")
    assert any("X_9.2" in chunk["text"] for chunk in found)


def test_table_row_comes_with_its_table(tmp_path, monkeypatch):
    _fresh(tmp_path, monkeypatch)
    folder = tmp_path / "docs"
    folder.mkdir()
    document = Document()
    document.add_heading("Склад", level=1)
    document.add_paragraph("Таблица 1 – Помещения склада")
    rows = [("Помещение", "Площадь"), ("Холодильная камера", "40 м²"), ("Сушильная", "25 м²"), ("Архив", "12 м²")]
    table = document.add_table(rows=len(rows), cols=2)
    for row, (key, value) in zip(table.rows, rows):
        row.cells[0].text, row.cells[1].text = key, value
    document.save(folder / "store.docx")
    store.upsert_chunks(ingest_dir(folder))
    found = store.search_chunks("Какая площадь у холодильной камеры?")
    text = next(chunk["text"] for chunk in found if "Холодильная камера" in chunk["text"])
    assert "Таблица 1 – Помещения склада" in text
    assert "Холодильная камера: 40 м²" in text
    assert "Архив: 12 м²" in text
    assert text.count("Таблица 1") == 1


def test_scope_names_one_file_and_stays_open_on_a_tie():
    assert store.scope({"vocabulary": {"only.txt": {"насос"}}}, "Что умеет насос Н-12?") is None
    assert store.scope(
        {"vocabulary": {"a.txt": {"альфа"}, "b.txt": {"альфа"}}},
        "Как устроена система АЛЬФА?",
    ) is None
    assert store.scope(
        {"vocabulary": {"a.txt": {"альфа"}, "b.txt": {"бета"}}},
        "Как устроена система БЕТА?",
    ) == ["b.txt"]
    assert store.scope(
        {"vocabulary": {"a.txt": {"альфа"}, "b.txt": {"бета"}}},
        "Как устроена система ГАММА?",
    ) is None
    numbered = {"vocabulary": {"STD_1.2-99.txt": set(), "STD_1.3-99.txt": set()}}
    assert store.scope(numbered, "Что указано в STD 1.2-99?") == ["STD_1.2-99.txt"]


def test_focus_drops_the_document_name_when_words_remain():
    scoped = store._focus("Какое давление у системы БЕТА сейчас", ["b.txt"])
    assert scoped == "Какое давление у системы сейчас"
    assert store._focus("БЕТА?", ["b.txt"]) == "БЕТА?"
    assert store._focus("Какое давление у системы БЕТА", None) == "Какое давление у системы БЕТА"


def test_table_context_keeps_overview_and_near_rows_inside_the_budget(monkeypatch):
    monkeypatch.setattr(ingest, "count_tokens", lambda text: max(1, len((text or "").split())))
    small = {
        "texts": ["цель", "сосед", "далеко"],
        "metas": [
            {"head": "шапка", "kind": "table", "document_name": "t.docx", "group": "t1", "chunk_index": 0},
            {"head": "", "kind": "table", "document_name": "t.docx", "group": "t1", "chunk_index": 1},
            {"head": "", "kind": "table", "document_name": "t.docx", "group": "t1", "chunk_index": 2},
        ],
        "groups": {("t.docx", "t1"): [0, 1, 2]},
    }
    monkeypatch.setattr(store, "_context", lambda: 3)
    near = store._render(small, [0])
    assert "цель" in near and "сосед" in near and "шапка" in near
    assert "далеко" not in near
    wide = {
        "texts": ["цель", "сосед", "один два три четыре пять шесть"],
        "metas": [
            {"head": "шапка", "kind": "table", "document_name": "t.docx", "group": "t1", "chunk_index": 0},
            {"head": "", "kind": "table", "document_name": "t.docx", "group": "t1", "chunk_index": 1},
            {"head": "", "kind": "overview", "document_name": "t.docx", "group": "t1", "chunk_index": 2},
        ],
        "groups": {("t.docx", "t1"): [0, 1, 2]},
    }
    monkeypatch.setattr(store, "_context", lambda: 4)
    forced = store._render(wide, [0])
    assert "один два три четыре пять шесть" in forced
    assert "цель" in forced
    assert "сосед" not in forced


def test_contexts_collapse_one_table_into_a_single_slot(monkeypatch):
    monkeypatch.setattr(ingest, "count_tokens", lambda text: 1)
    monkeypatch.setattr(store, "_context", lambda: 100)
    index = {
        "ids": ["a0", "a1", "a2", "b0"],
        "texts": ["A0", "A1", "список значений", "B0"],
        "metas": [
            {"document_name": "a.docx", "group": "t1", "kind": "table", "chunk_index": 0, "section": "S", "head": ""},
            {"document_name": "a.docx", "group": "t1", "kind": "table", "chunk_index": 1, "section": "S", "head": ""},
            {"document_name": "a.docx", "group": "t1", "kind": "overview", "chunk_index": 2, "section": "S", "head": ""},
            {"document_name": "b.txt", "group": "", "kind": "text", "chunk_index": 0, "section": "", "head": ""},
        ],
        "groups": {("a.docx", "t1"): [0, 1, 2]},
        "position": {"a0": 0, "a1": 1, "a2": 2, "b0": 3},
    }
    found = store._contexts(index, ["a0", "a1", "a2", "b0"], 2)
    assert [item["document_name"] for item in found] == ["a.docx", "b.txt"]
    assert found[0]["kind"] == "overview"
    assert "A0" in found[0]["text"] and "A1" in found[0]["text"] and "список значений" in found[0]["text"]
    assert found[1]["text"] == "B0"


def test_fuse_ranks_a_shared_hit_above_a_single_list(monkeypatch):
    monkeypatch.setattr(store, "_rrf", lambda: 5)
    assert store._fuse(["b", "a"], ["a"]) == ["a", "b"]


def test_blank_route_values_are_not_filters(monkeypatch):
    monkeypatch.setattr(settings, "routing", lambda: {"where": {"collection": "", "kind": None, "tags": [], "keep": "alpha"}})
    assert settings.where() == {"keep": "alpha"}
    monkeypatch.setattr(store, "route_where", lambda: {})
    assert store._meta_ok({"collection": "beta"}) is True
    assert store._chroma_where(None) is None
    assert store._chroma_where(["a.txt"]) == {"document_name": {"$in": ["a.txt"]}}
    monkeypatch.setattr(store, "route_where", lambda: {"collection": "alpha"})
    assert store._meta_ok({"collection": "alpha"}) is True
    assert store._meta_ok({}) is False
    assert store._chroma_where(None) == {"collection": "alpha"}
    assert store._chroma_where(["a.txt"]) == {
        "$and": [{"collection": "alpha"}, {"document_name": {"$in": ["a.txt"]}}],
    }


def test_lexical_hides_another_collection(monkeypatch):
    monkeypatch.setattr(store, "route_where", lambda: {"collection": "alpha"})
    monkeypatch.setattr(store, "_candidates", lambda: 10)
    terms = store._terms("насос")
    count = Counter({term: 1 for term in terms})
    index = {
        "counts": [Counter(count), Counter(count)],
        "lengths": [sum(count.values()), sum(count.values())],
        "metas": [
            {"document_name": "a.txt", "collection": "alpha"},
            {"document_name": "b.txt", "collection": "beta"},
        ],
        "frequency": {"a.txt": Counter(count), "b.txt": Counter(count)},
        "size": Counter({"a.txt": sum(count.values()), "b.txt": sum(count.values())}),
        "chunks": Counter({"a.txt": 1, "b.txt": 1}),
        "ids": ["a", "b"],
    }
    assert store._lexical(index, "насос", None) == ["a"]
    assert store._lexical(index, "насос", ["a.txt", "b.txt"]) == ["a"]
    monkeypatch.setattr(store, "route_where", lambda: {})
    assert store._lexical(index, "насос", ["b.txt"]) == ["b"]
    assert store._lexical(index, "склад", ["a.txt"]) == []
