from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ingest import TOKENS, VERSION, _docx_blocks, count_tokens, ingest_dir, load_text, split_chunks
import ingest
from ingest import TOKENS, VERSION, count_tokens, ingest_dir, load_text, split_chunks


def _toc_style(document):
    try:
        return document.styles["toc 1"]
    except KeyError:
        return document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)


def test_txt_and_md(tmp_path):
    (tmp_path / "a.txt").write_text("Цех имеет номер 4.", encoding="utf-8")
    (tmp_path / "b.md").write_text("# Насос\nНасос Н-12.", encoding="utf-8")
    chunks = {chunk["document_name"]: chunk for chunk in ingest_dir(tmp_path)}
    assert set(chunks) == {"a.txt", "b.md"}
    assert chunks["b.md"]["section"] == "Насос"
    assert chunks["b.md"]["text"] == "Насос\nНасос Н-12."


def test_docx_keeps_order_sections_and_skips_toc(tmp_path):
    path = tmp_path / "c.docx"
    document = Document()
    document.add_paragraph("Склад 7", style=_toc_style(document))
    document.add_heading("Склад", level=1)
    document.add_paragraph("Температура склада равна 18 градусам.")
    document.add_paragraph("Таблица 1 – Состав склада")
    table = document.add_table(rows=3, cols=2)
    for row, (key, value) in zip(table.rows, [("Показатель", "Значение"), ("Площадь", "120 м²"), ("Этаж", "2")]):
        row.cells[0].text, row.cells[1].text = key, value
    document.add_paragraph("Склад закрыт по воскресеньям.")
    document.save(path)
    chunks = ingest_dir(tmp_path)
    texts = [chunk["text"] for chunk in chunks]
    assert not any("Склад 7" in text for text in texts)
    rows = [chunk for chunk in chunks if chunk["kind"] == "table"]
    assert [chunk["section"] for chunk in rows] == ["Склад", "Склад"]
    assert all("Таблица 1 – Состав склада" in chunk["text"] for chunk in rows)
    assert "Площадь: 120 м²" in rows[0]["text"]
    order = [text for text in texts if "18 градусам" in text or "Площадь" in text or "воскресеньям" in text]
    assert "18 градусам" in order[0] and "воскресеньям" in order[-1]
    intro = next(chunk for chunk in chunks if chunk["kind"] == "intro")
    assert "18 градусам" in intro["text"] and intro["group"] == rows[0]["group"]
    assert not any(chunk["kind"] == "overview" for chunk in chunks)
    assert "18" in load_text(path)


def test_overview_list_keeps_every_value(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", lambda text: max(1, len(text.split())))
    from ingest import _word_tokens, chunks_from_blocks

    _word_tokens.cache_clear()
    values = [f"Наименование характеристики номер {number} очень длинное" for number in range(1, 41)]
    blocks = [{
        "kind": "table",
        "caption": "Таблица 1 – Перечень",
        "labels": ["Статус", "Описание"],
        "width": 2,
        "rows": [[("Статус", "Да"), ("Описание", value)] for value in values],
    }]
    chunks = chunks_from_blocks(blocks, "list.txt")
    overview = next(chunk["text"] for chunk in chunks if chunk["kind"] == "overview")
    assert "Разных значений в столбце «Описание»: 40." in overview
    assert all(value in overview for value in values)


def test_large_table_gets_overview(tmp_path):
    document = Document()
    table = document.add_table(rows=41, cols=3)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text, table.rows[0].cells[2].text = "Код", "Описание", "Статус"
    for number, row in enumerate(table.rows[1:], 1):
        row.cells[0].text = f"K_{number}"
        row.cells[1].text = f"Описание позиции номер {number} в длинном перечне склада"
        row.cells[2].text = "Да"
    document.save(tmp_path / "big.docx")
    chunks = ingest_dir(tmp_path)
    overview = next(chunk for chunk in chunks if chunk["kind"] == "overview")
    assert "Разных значений в столбце «Код»: 40" in overview["text"]
    assert len({chunk["group"] for chunk in chunks if chunk["kind"] in ("overview", "table")}) == 1


def test_merged_cells_are_not_repeated(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "ИД"
    table.rows[0].cells[1].merge(table.rows[0].cells[2]).text = "Описание"
    table.rows[1].cells[0].text = "A_1"
    table.rows[1].cells[1].text = "Запрос"
    table.rows[1].cells[2].text = "Передача данных"
    document.save(tmp_path / "m.docx")
    row = next(chunk["text"] for chunk in ingest_dir(tmp_path) if chunk["kind"] == "table")
    assert row.count("Описание") == 1
    assert "ИД: A_1; Описание: Запрос; Передача данных" in row


def test_table_row_that_starts_late_keeps_its_columns(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from ingest import _docx_blocks

    document = Document()
    table = document.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Показатель", "Минимум", "Максимум")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("", "6 бар", "10 бар")):
        cell.text = text
    data = table.rows[1]._tr
    data.remove(data.tc_lst[0])
    mark = OxmlElement("w:gridBefore")
    mark.set(qn("w:val"), "1")
    data.get_or_add_trPr().append(mark)
    document.save(tmp_path / "range.docx")
    block = _docx_blocks(tmp_path / "range.docx")[0]
    assert block["rows"] == [[("Минимум", "6 бар"), ("Максимум", "10 бар")]]


def test_chunks_fit_tokens_and_overlap():
    text = " ".join(f"слово{i}" for i in range(900))
    chunks = split_chunks(text, "long.txt")
    assert len(chunks) > 1
    assert all(count_tokens(chunk["text"]) <= TOKENS for chunk in chunks)
    assert len({chunk["checksum"] for chunk in chunks}) == 1
    assert all(chunk["version"] == VERSION for chunk in chunks)
    first, second = chunks[0]["text"].split(), chunks[1]["text"].split()
    assert first[-1] in second[:40]


def test_long_line_splits_after_a_sentence(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", lambda text: max(1, len(str(text).split())))
    from ingest import _cut, _word_tokens

    _word_tokens.cache_clear()
    lead = " ".join(f"слово{number}" for number in range(50)) + "."
    mid = " ".join(f"хвост{number}" for number in range(40))
    tail = "В разделе Назначение должны быть указаны сведения о назначении программы."
    pieces = _cut((lead + " " + mid + " " + tail).split(), 80)
    assert pieces[0].endswith(".")
    assert "Назначение" not in pieces[0]
    assert any("Назначение" in piece for piece in pieces[1:])
    _word_tokens.cache_clear()


def test_docx_does_not_join_a_hyphen_to_a_new_sentence(tmp_path):
    from ingest import _docx_blocks

    document = Document()
    document.add_paragraph("Строка заканчивается дефисом-")
    document.add_paragraph("Новое предложение начинается с большой буквы.")
    path = tmp_path / "hyphen.docx"
    document.save(path)
    texts = [block["text"] for block in _docx_blocks(path)]
    assert "Строка заканчивается дефисом-" in texts
    assert "Новое предложение начинается с большой буквы." in texts
    assert all("дефисомНовое" not in text for text in texts)


def test_docx_joins_hyphen_line_break(tmp_path):
    document = Document()
    document.add_paragraph("Требования к содержанию доку-")
    document.add_paragraph("мента приведены в приложении.")
    document.save(tmp_path / "hyphen.docx")
    text = " ".join(chunk["text"] for chunk in ingest_dir(tmp_path))
    assert "документа" in text
    assert "доку-" not in text


def test_docx_keeps_hard_hyphen_before_a_number(tmp_path):
    document = Document()
    document.add_paragraph("Требования к содержанию доку-")
    document.add_paragraph("мента приведены в приложении.")
    document.add_paragraph("Обозначение приведено в ГОСТ 19.201-")
    document.add_paragraph("78 и действует далее.")
    document.add_paragraph("Рабочий диапазон 6-")
    document.add_paragraph("10 бар указан в паспорте.")
    document.add_paragraph("Норма ISO-")
    document.add_paragraph("9001 указана в договоре.")
    document.add_paragraph("Смысл значения -")
    document.add_paragraph("ниже предела.")
    document.save(tmp_path / "hyphen.docx")
    from ingest import _docx_blocks

    text = "\n".join(block["text"] for block in _docx_blocks(tmp_path / "hyphen.docx"))
    assert "документа" in text
    assert "доку-" not in text
    assert "ГОСТ 19.201-\n78" in text
    assert "19.20178" not in text
    assert "6-\n10" in text
    assert "610" not in text
    assert "ISO9001" not in text
    assert "ISO-" in text
    assert "значения -" in text
    assert "значения ниже" not in text


def test_nested_table_keeps_the_cell_value(tmp_path):
    document = Document()
    document.add_paragraph("Паспорт насоса.")
    outer = document.add_table(rows=2, cols=2)
    outer.rows[0].cells[0].text = "Показатель"
    outer.rows[0].cells[1].text = "Значение"
    outer.rows[1].cells[0].text = "Давление"
    value = outer.rows[1].cells[1]
    value.text = "не ниже"
    nested = value.add_table(rows=2, cols=2)
    nested.rows[0].cells[0].text = "Режим"
    nested.rows[0].cells[1].text = "Бар"
    nested.rows[1].cells[0].text = "Рабочий"
    nested.rows[1].cells[1].text = "10 бар"
    document.save(tmp_path / "nested.docx")
    blocks = _docx_blocks(tmp_path / "nested.docx")
    table = next(block for block in blocks if block["kind"] == "table")
    text = " ".join(value for row in table["rows"] for _, value in row)
    assert "Паспорт насоса." in blocks[0]["text"]
    assert "не ниже" in text
    assert "Рабочий: 10 бар" in text
    assert table["rows"][0][0] == ("Показатель", "Давление")


def test_paragraph_with_section_number_is_kept():
    chunks = split_chunks("Согласно разделу 3.2 насос имеет давление 6 бар.", "pump.txt")
    assert "3.2" in chunks[0]["text"]


def test_unreadable_file_does_not_block_other_documents(tmp_path, monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", lambda text: max(1, len(str(text).split())))
    from ingest import _word_tokens

    _word_tokens.cache_clear()

    def converter_missing(path):
        raise FileNotFoundError(path)

    monkeypatch.setattr("ingest._doc_blocks", converter_missing)
    (tmp_path / "good.txt").write_text("Цех имеет номер 4.", encoding="utf-8")
    (tmp_path / "bad.txt").write_bytes(b"\xff")
    (tmp_path / "broken.docx").write_bytes(b"not a docx")
    (tmp_path / "legacy.doc").write_bytes(b"doc")
    chunks = ingest_dir(tmp_path)
    assert {chunk["document_name"] for chunk in chunks} == {"good.txt"}
    assert any("4" in chunk["text"] for chunk in chunks)
    _word_tokens.cache_clear()


def _word_run(text):
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    return run


def _word_mark(tag, *children):
    node = OxmlElement(tag)
    node.set(qn("w:id"), "1")
    node.set(qn("w:author"), "Ed")
    node.set(qn("w:date"), "2026-10-31T00:00:00Z")
    for child in children:
        node.append(child)
    return node


def _word_sdt(*children):
    sdt = OxmlElement("w:sdt")
    sdt.append(OxmlElement("w:sdtPr"))
    content = OxmlElement("w:sdtContent")
    for child in children:
        content.append(child)
    sdt.append(content)
    return sdt


def _save(document, path):
    document.save(path)
    return path


def test_docx_keeps_an_inserted_value_and_drops_the_deleted_one(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Давление равно ")
    deleted = _word_mark("w:del", _word_run("6 бар"))
    # A deletion stores its text in delText. A normal run would still be visible.
    deleted.find(qn("w:r")).find(qn("w:t")).tag = qn("w:delText")
    paragraph._p.append(deleted)
    paragraph._p.append(_word_mark("w:ins", _word_run("10 бар")))
    _save(document, tmp_path / "revised.docx")
    text = " ".join(block["text"] for block in _docx_blocks(tmp_path / "revised.docx"))
    assert "10 бар" in text
    assert "6 бар" not in text


def test_docx_keeps_a_value_inside_a_content_control(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Давление равно ")
    paragraph._p.append(_word_sdt(_word_run("10 бар")))
    _save(document, tmp_path / "field.docx")
    text = " ".join(block["text"] for block in _docx_blocks(tmp_path / "field.docx"))
    assert "10 бар" in text


def test_docx_keeps_a_paragraph_wrapped_in_a_content_control(tmp_path):
    document = Document()
    body = document.element.body
    paragraph = OxmlElement("w:p")
    paragraph.append(_word_run("Рабочее давление равно 10 бар."))
    body.insert(len(body) - 1, _word_sdt(paragraph))
    _save(document, tmp_path / "block-field.docx")
    text = " ".join(block["text"] for block in _docx_blocks(tmp_path / "block-field.docx"))
    assert "Рабочее давление равно 10 бар." in text


def test_docx_keeps_an_inserted_table_cell(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Показатель"
    table.rows[0].cells[1].text = "Значение"
    table.rows[1].cells[0].text = "Давление"
    table.rows[1].cells[1].paragraphs[0]._p.append(_word_mark("w:ins", _word_run("10 бар")))
    _save(document, tmp_path / "table.docx")
    table = next(block for block in _docx_blocks(tmp_path / "table.docx") if block["kind"] == "table")
    values = [value for row in table["rows"] for _, value in row]
    assert "Давление" in values
    assert "10 бар" in values


def test_docx_keeps_a_nested_control_inside_an_insertion(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Код ")
    paragraph._p.append(_word_mark("w:ins", _word_sdt(_word_run("Н-12"))))
    _save(document, tmp_path / "nested.docx")
    text = " ".join(block["text"] for block in _docx_blocks(tmp_path / "nested.docx"))
    assert "Н-12" in text


def test_docx_keeps_text_when_a_revision_only_marks_formatting(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("Давление равно 6 бар.")
    mark = OxmlElement("w:ins")
    mark.set(qn("w:id"), "7")
    mark.set(qn("w:author"), "Ed")
    mark.set(qn("w:date"), "2026-10-31T00:00:00Z")
    paragraph.runs[0]._r.get_or_add_rPr().append(mark)
    _save(document, tmp_path / "format.docx")
    text = " ".join(block["text"] for block in _docx_blocks(tmp_path / "format.docx"))
    assert text == "Давление равно 6 бар."


def test_docx_revision_wrapper_leaves_a_plain_paragraph(tmp_path):
    document = Document()
    document.add_paragraph("Требования к содержанию доку-")
    document.add_paragraph("мента приведены в приложении.")
    _save(document, tmp_path / "plain.docx")
    blocks = _docx_blocks(tmp_path / "plain.docx")
    assert blocks[0]["text"] == "Требования к содержанию документа приведены в приложении."


def test_doc_via_libreoffice(tmp_path):
    soffice = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
    if not soffice.exists():
        pytest.skip("LibreOffice не установлен")
    source = tmp_path / "d.docx"
    document = Document()
    document.add_paragraph("Год выпуска равен 2019.")
    document.save(source)
    import subprocess
    subprocess.run([str(soffice), "--headless", "--convert-to", "doc", "--outdir", str(tmp_path), str(source)], check=True)
    source.unlink()
    assert "2019" in load_text(tmp_path / "d.doc")
    assert sorted(path.name for path in tmp_path.iterdir()) == ["d.doc"]


def test_numbered_line_becomes_section(tmp_path):
    path = tmp_path / "n.docx"
    document = Document()
    document.add_paragraph("3 Перечень испытаний")
    document.add_paragraph("Опытная эксплуатация проверяет работу.")
    document.add_paragraph("2.1. В разделе Введение указывают наименование программы и объекта, в котором используют программу.")
    document.save(path)
    chunks = ingest_dir(tmp_path)
    assert chunks[0]["section"] == "3 Перечень испытаний"
    assert "Опытная эксплуатация" in chunks[0]["text"]
    assert all(chunk["section"] != "2.1. В разделе Введение указывают наименование программы и объекта, в котором используют программу." for chunk in chunks)

def test_docx_does_not_join_a_capitalized_line(tmp_path):
    from ingest import _docx_blocks

    document = Document()
    document.add_paragraph("Требования к содержанию доку-")
    document.add_paragraph("Мента приведены отдельно.")
    path = tmp_path / "split.docx"
    document.save(path)
    texts = [block["text"] for block in _docx_blocks(path)]
    assert texts == [
        "Требования к содержанию доку-",
        "Мента приведены отдельно.",
    ]

def _token_count(text):
    return max(1, len(str(text or "").split()))

def test_long_line_does_not_split_a_hyphenated_word(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", _token_count)
    from ingest import _cut, _word_tokens

    _word_tokens.cache_clear()
    try:
        words = ["доку-мента"] + [f"слово{number}" for number in range(30)]
        pieces = _cut(words, 10)
        assert pieces[0].startswith("доку-мента")
        assert all(not piece.startswith("мента") for piece in pieces)
    finally:
        _word_tokens.cache_clear()

def test_heading_path_resets_and_does_not_repeat_the_caption(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", _token_count)
    monkeypatch.setattr(ingest, "runtime", lambda: {"tokens": 80, "overlap": 0, "intro": 120})
    from ingest import _word_tokens

    _word_tokens.cache_clear()
    try:
        blocks = [
            {"kind": "heading", "level": 1, "text": "Общее"},
            {"kind": "heading", "level": 2, "text": "Назначение"},
            {"kind": "p", "text": "Текст назначения."},
            {"kind": "heading", "level": 2, "text": "Состав"},
            {"kind": "p", "text": "Текст состава."},
            {
                "kind": "table",
                "caption": "Состав",
                "labels": ["Код", "Имя"],
                "width": 2,
                "rows": [[("Код", "A"), ("Имя", "Насос")]],
            },
            {"kind": "heading", "level": 1, "text": "Требования"},
            {"kind": "p", "text": "Текст требований."},
        ]
        chunks = ingest.chunks_from_blocks(blocks, "doc.md")
        assert [chunk["section"] for chunk in chunks] == [
            "Общее > Назначение",
            "Общее > Состав",
            "Общее > Состав",
            "Требования",
        ]
        assert chunks[0]["kind"] == "text" and "Текст назначения." in chunks[0]["text"]
        assert chunks[1]["kind"] == "intro" and "Текст состава." in chunks[1]["text"]
        assert chunks[2]["kind"] == "table"
        assert chunks[2]["text"].splitlines()[0] == "Общее > Состав"
        assert chunks[2]["text"].count("Состав") == 1
        assert "A: Насос" in chunks[2]["text"]
        assert "Текст требований." in chunks[3]["text"]
    finally:
        _word_tokens.cache_clear()

def test_long_paragraph_before_a_table_stays_text(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", _token_count)
    monkeypatch.setattr(ingest, "runtime", lambda: {"tokens": 80, "overlap": 0, "intro": 5})
    from ingest import _word_tokens

    _word_tokens.cache_clear()
    try:
        words = " ".join(f"слово{number}" for number in range(12))
        blocks = [
            {"kind": "p", "text": words},
            {
                "kind": "table",
                "caption": "Таблица 1",
                "labels": ["Код", "Имя"],
                "width": 2,
                "rows": [[("Код", "A"), ("Имя", "Насос")]],
            },
        ]
        chunks = ingest.chunks_from_blocks(blocks, "t.txt")
        assert [chunk["kind"] for chunk in chunks] == ["text", "table"]
        assert "слово0" in chunks[0]["text"] and "слово11" in chunks[0]["text"]
        assert "A: Насос" in chunks[1]["text"]
    finally:
        _word_tokens.cache_clear()

def test_wide_table_row_keeps_the_lead_columns(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", _token_count)
    monkeypatch.setattr(ingest, "runtime", lambda: {"tokens": 40, "overlap": 0, "intro": 5})
    from ingest import _word_tokens

    _word_tokens.cache_clear()
    try:
        def column(prefix):
            return " ".join(f"{prefix}{number}" for number in range(40))

        blocks = [{
            "kind": "table",
            "caption": "Таблица 1",
            "labels": ["Код", "Имя", "Цех", "Примечание"],
            "width": 4,
            "rows": [[
                ("Код", "A1"),
                ("Имя", column("имя")),
                ("Цех", column("цех")),
                ("Примечание", column("хвост")),
            ]],
        }]
        tables = [chunk["text"] for chunk in ingest.chunks_from_blocks(blocks, "wide.txt") if chunk["kind"] == "table"]
        assert len(tables) == 3
        assert all("Код: A1" in text for text in tables)
        assert "имя0" in tables[0] and "цех0" in tables[1] and "хвост39" in tables[2]
    finally:
        _word_tokens.cache_clear()

def test_docx_does_not_join_a_new_sentence_a_heading_or_text_after_a_table(tmp_path):
    document = Document()
    document.add_paragraph("Требования к содержанию доку-")
    document.add_paragraph("Мента приведены отдельно.")
    document.add_paragraph("Строка заканчивается дефи-")
    document.add_heading("Новый раздел", level=1)
    document.add_paragraph("Хвост перед табли-")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Код"
    table.rows[0].cells[1].text = "Имя"
    table.rows[1].cells[0].text = "A"
    table.rows[1].cells[1].text = "Насос"
    document.add_paragraph("цей не склеивается.")
    path = tmp_path / "bounds.docx"
    document.save(path)
    texts = [block.get("text", "") for block in ingest._docx_blocks(path)]
    assert "Требования к содержанию доку-" in texts
    assert "Мента приведены отдельно." in texts
    assert "Строка заканчивается дефи-" in texts
    assert "Новый раздел" in texts
    assert "Хвост перед табли-" in texts
    assert "цей не склеивается." in texts
    joined = " ".join(texts)
    assert "документа" not in joined
    assert "дефи-Новый" not in joined
    assert "таблицей" not in joined

def test_ingest_dir_indexes_only_supported_files(tmp_path, monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", _token_count)
    from ingest import _word_tokens

    _word_tokens.cache_clear()
    try:
        (tmp_path / "a.txt").write_text("Первая строка.", encoding="utf-8")
        (tmp_path / "skip.pdf").write_bytes(b"\xff\xfe")
        nested = tmp_path / "nested"
        nested.mkdir()
        (nested / "b.txt").write_text("Вложенный текст.", encoding="utf-8")
        chunks = ingest_dir(tmp_path)
        assert [chunk["document_name"] for chunk in chunks] == ["a.txt"]
        assert "Первая строка." in chunks[0]["text"]
    finally:
        _word_tokens.cache_clear()

def test_docx_keeps_a_hyphen_before_a_heading_or_a_capital(tmp_path):
    document = Document()
    document.add_paragraph("Конец строки-")
    document.add_paragraph("Новое предложение остаётся отдельным.")
    document.add_paragraph("префикс-")
    document.add_heading("Заголовок не склеивается", level=1)
    document.add_heading("Раздел-", level=2)
    document.add_paragraph("не часть заголовка.")
    document.add_paragraph("слово-")
    document.add_paragraph("продолжение в том же слове.")
    document.save(tmp_path / "breaks.docx")
    from ingest import _docx_blocks

    texts = [(block["kind"], block["text"]) for block in _docx_blocks(tmp_path / "breaks.docx")]
    assert texts == [
        ("p", "Конец строки-"),
        ("p", "Новое предложение остаётся отдельным."),
        ("p", "префикс-"),
        ("heading", "Заголовок не склеивается"),
        ("heading", "Раздел-"),
        ("p", "не часть заголовка."),
        ("p", "словопродолжение в том же слове."),
    ]

def test_long_line_ignores_a_sentence_end_in_the_first_half(monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", lambda text: max(1, len(str(text).split())))
    from ingest import _cut, _word_tokens

    _word_tokens.cache_clear()
    early = ["Короткое."] + [f"слово{number}" for number in range(20)]
    pieces = _cut(early + ["Хвост", "готов?"], 20)
    assert pieces[0].startswith("Короткое.")
    assert pieces[0].endswith("слово18")
    assert "Хвост" not in pieces[0]
    head = [f"слово{number}" for number in range(12)]
    tail = ["граница?"] + [f"хвост{number}" for number in range(12)]
    pieces = _cut(head + tail, 20)
    assert pieces[0].endswith("граница?")
    assert "хвост0" not in pieces[0]
    _word_tokens.cache_clear()


def test_docx_keeps_text_moved_by_a_revision(tmp_path):
    document = Document()
    body = document.element.body
    old = OxmlElement("w:p")
    old.append(_word_run("Давление равно 6 бар."))
    new = OxmlElement("w:p")
    new.append(_word_run("Давление равно 10 бар."))
    body.insert(len(body) - 1, _word_mark("w:moveFrom", old))
    body.insert(len(body) - 1, _word_mark("w:moveTo", new))
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Показатель"
    table.rows[0].cells[1].text = "Значение"
    table.rows[1].cells[0].text = "Давление"
    value = table.rows[1].cells[1].paragraphs[0]._p
    value.append(_word_mark("w:moveFrom", _word_run("6 бар")))
    value.append(_word_mark("w:moveTo", _word_sdt(_word_run("10 бар"))))
    _save(document, tmp_path / "moved.docx")
    blocks = _docx_blocks(tmp_path / "moved.docx")
    texts = []
    for block in blocks:
        if block["kind"] == "table":
            texts.extend(value for row in block["rows"] for _, value in row)
        else:
            texts.append(block["text"])
    text = "\n".join(texts)
    assert "Давление равно 10 бар." in text
    assert "Давление равно 6 бар." not in text
    assert "10 бар" in text
    assert "6 бар" not in text


def test_missing_parser_aborts_instead_of_skipping_the_file(tmp_path, monkeypatch):
    monkeypatch.setattr("ingest.count_tokens", lambda text: max(1, len(str(text or "").split())))
    from ingest import _word_tokens

    _word_tokens.cache_clear()

    def missing(_path):
        raise ImportError("python-docx")

    monkeypatch.setattr("ingest._docx_blocks", missing)
    (tmp_path / "a.txt").write_text("Цех имеет номер 4.", encoding="utf-8")
    (tmp_path / "z.docx").write_bytes(b"PK")
    with pytest.raises(ImportError, match="python-docx"):
        ingest_dir(tmp_path)
    _word_tokens.cache_clear()
